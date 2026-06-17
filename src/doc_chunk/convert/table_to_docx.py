"""Render physical table grids back into Word documents."""

from __future__ import annotations

from typing import Any

from docx.document import Document
from docx.table import Table

from doc_chunk.models.table_model import TableGridRow, TableSidecar


def _parse_grid_rows(grid: dict[str, Any]) -> list[TableGridRow]:
    return [TableGridRow.model_validate(row) for row in grid.get("rows", [])]


def _place_physical_grid(
    table: Table,
    physical_rows: list[TableGridRow],
    grid_width: int,
) -> list[list[tuple[int, int]]]:
    """Write cells and merges; return anchor (row, col) per physical cell."""
    row_count = len(physical_rows)
    occupied = [[False] * grid_width for _ in range(row_count)]
    positions: list[list[tuple[int, int]]] = []

    for r_idx, row in enumerate(physical_rows):
        row_positions: list[tuple[int, int]] = []
        col_idx = 0
        for cell in row.cells:
            while col_idx < grid_width and occupied[r_idx][col_idx]:
                col_idx += 1
            if col_idx >= grid_width:
                break

            end_col = min(col_idx + cell.colspan - 1, grid_width - 1)
            end_row = min(r_idx + cell.rowspan - 1, row_count - 1)

            anchor = table.cell(r_idx, col_idx)
            anchor.text = cell.text
            if end_col > col_idx or end_row > r_idx:
                anchor.merge(table.cell(end_row, end_col))

            for rr in range(r_idx, end_row + 1):
                for cc in range(col_idx, end_col + 1):
                    occupied[rr][cc] = True

            row_positions.append((r_idx, col_idx))
            col_idx = end_col + 1
        positions.append(row_positions)

    return positions


def _apply_records_to_table(
    table: Table,
    physical_rows: list[TableGridRow],
    positions: list[list[tuple[int, int]]],
    records: list[dict[str, str]],
    *,
    layout_type: str,
    logical_rows: list[list[str]],
) -> None:
    if not records:
        return

    if layout_type == "personnel_dual_row":
        for rec_idx, record in enumerate(records):
            base = rec_idx * 4
            for header_offset, data_offset in ((0, 1), (2, 3)):
                header_idx = base + header_offset
                data_idx = base + data_offset
                if data_idx >= len(physical_rows):
                    return
                for cell_i, header_cell in enumerate(physical_rows[header_idx].cells):
                    key = header_cell.text.strip()
                    if not key or key not in record or cell_i >= len(positions[data_idx]):
                        continue
                    row, col = positions[data_idx][cell_i]
                    table.cell(row, col).text = record[key]
        return

    if layout_type == "simple" and logical_rows:
        headers = [h.strip() for h in logical_rows[0]]
        for rec_idx, record in enumerate(records):
            data_idx = rec_idx + 1
            if data_idx >= len(positions):
                break
            for col_i, header in enumerate(headers):
                if not header or header not in record or col_i >= len(positions[data_idx]):
                    continue
                row, col = positions[data_idx][col_i]
                table.cell(row, col).text = record[header]
        return

    if layout_type == "key_value" and records:
        merged = records[0]
        for row_idx, row in enumerate(physical_rows):
            if row_idx >= len(logical_rows) or len(positions[row_idx]) < 2:
                continue
            key = logical_rows[row_idx][0].strip() if logical_rows[row_idx] else ""
            if not key or key not in merged:
                continue
            row_pos, col_pos = positions[row_idx][1]
            table.cell(row_pos, col_pos).text = merged[key]
        return

    # fallback: match header cells anywhere in grid
    for record in records:
        for row_idx, row in enumerate(physical_rows):
            for cell_i, cell in enumerate(row.cells):
                key = cell.text.strip()
                if key in record and cell_i < len(positions[row_idx]):
                    # only update if next row exists and same column slot
                    if row_idx + 1 < len(positions) and cell_i < len(positions[row_idx + 1]):
                        r, c = positions[row_idx + 1][cell_i]
                        table.cell(r, c).text = record[key]


def render_table_to_docx(
    document: Document,
    grid: dict[str, Any],
    *,
    grid_width: int,
    records: list[dict[str, str]] | None = None,
    layout_type: str | None = None,
    logical_rows: list[list[str]] | None = None,
    record_groups: list[list[int]] | None = None,
) -> Table:
    del record_groups  # reserved for future multi-block layouts
    physical_rows = _parse_grid_rows(grid)
    if not physical_rows:
        raise ValueError("grid has no rows")
    if grid_width <= 0:
        raise ValueError("grid_width must be positive")

    table = document.add_table(rows=len(physical_rows), cols=grid_width)
    positions = _place_physical_grid(table, physical_rows, grid_width)

    if records:
        _apply_records_to_table(
            table,
            physical_rows,
            positions,
            records,
            layout_type=layout_type or "fallback",
            logical_rows=logical_rows or [],
        )

    return table


def render_sidecar_to_docx(
    document: Document,
    sidecar: TableSidecar,
    *,
    records: list[dict[str, str]] | None = None,
) -> Table:
    return render_table_to_docx(
        document,
        sidecar.grid,
        grid_width=sidecar.grid_width,
        records=records if records is not None else sidecar.records,
        layout_type=sidecar.layout_type,
        logical_rows=sidecar.logical_rows,
        record_groups=sidecar.record_groups,
    )
