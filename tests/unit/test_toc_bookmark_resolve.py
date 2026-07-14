from __future__ import annotations

import io
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from doc_chunk.api import extract_file
from doc_chunk.models.content_block import ContentBlocksFile
from doc_chunk.models.outline import Anchor, OutlineNode, OutlineTree
from doc_chunk.outline.toc_bookmark import apply_toc_bookmark_anchors
from doc_chunk.outline.toc_docx import extract_docx_toc_outline

_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS = {"w": _WORD_NS}


def _add_bookmark(paragraph, bookmark_name: str, bookmark_id: int = 0) -> None:
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), str(bookmark_id))
    bookmark_start.set(qn("w:name"), bookmark_name)
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, bookmark_start)
    paragraph._p.append(bookmark_end)


def _make_toc_para(style_id: str, texts: list[str], instr: str | None = None) -> etree._Element:
    p = etree.Element(f"{{{_WORD_NS}}}p")
    p_pr = etree.SubElement(p, f"{{{_WORD_NS}}}pPr")
    p_style = etree.SubElement(p_pr, f"{{{_WORD_NS}}}pStyle")
    p_style.set(f"{{{_WORD_NS}}}val", style_id)
    if instr:
        run = etree.SubElement(p, f"{{{_WORD_NS}}}r")
        instr_text = etree.SubElement(run, f"{{{_WORD_NS}}}instrText")
        instr_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        instr_text.text = instr
    for text in texts:
        run = etree.SubElement(p, f"{{{_WORD_NS}}}r")
        t = etree.SubElement(run, f"{{{_WORD_NS}}}t")
        t.text = text
    return p


def _docx_toc_with_bookmark(tmp_path: Path) -> Path:
    docx_path = tmp_path / "toc-bookmark.docx"
    doc = Document()
    heading = doc.add_heading("一、服务方案", level=1)
    _add_bookmark(heading, "_Toc1")
    doc.add_paragraph("正文")
    doc.save(docx_path)

    buf_out = io.BytesIO()
    with ZipFile(docx_path, "r") as zin, ZipFile(buf_out, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                body = root.find(f"{{{_WORD_NS}}}body")
                assert body is not None
                toc_field = _make_toc_para("16", ["目录"], ' TOC \\o "1-3" \\h \\u ')
                toc_entry = _make_toc_para("16", ["一、服务方案", "1"], " HYPERLINK \\l _Toc1 ")
                body.insert(0, toc_entry)
                body.insert(0, toc_field)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item.filename == "word/styles.xml":
                root = etree.fromstring(data)
                if not root.xpath('.//w:style[@w:styleId="16"]', namespaces=_NS):
                    style = etree.SubElement(root, f"{{{_WORD_NS}}}style")
                    style.set(f"{{{_WORD_NS}}}type", "paragraph")
                    style.set(f"{{{_WORD_NS}}}styleId", "16")
                    name = etree.SubElement(style, f"{{{_WORD_NS}}}name")
                    name.set(f"{{{_WORD_NS}}}val", "toc 1")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    docx_path.write_bytes(buf_out.getvalue())
    return docx_path


def test_apply_toc_bookmark_anchors_sets_body_block(tmp_path: Path) -> None:
    src = _docx_toc_with_bookmark(tmp_path)
    ws = tmp_path / "ws"
    extract_file(src, ws, overwrite=True)
    tree = extract_docx_toc_outline(src)
    assert tree is not None
    blocks = ContentBlocksFile.model_validate_json((ws / "content.blocks.json").read_text())
    content_md = (ws / "content.md").read_text(encoding="utf-8")
    enriched = apply_toc_bookmark_anchors(tree, src, blocks, content_md=content_md)
    node = next(n for n in enriched.nodes if "服务方案" in n.title and n.level == 1)
    assert node.anchor.block_index is not None
    block = next(b for b in blocks.blocks if b.block_index == node.anchor.block_index)
    preview = block.text_preview or content_md[block.char_start : block.char_end]
    assert "一、服务方案" in preview
    assert "服务方案1" not in (block.text_preview or "")


def test_apply_toc_bookmark_anchors_missing_bookmark_noop(tmp_path: Path) -> None:
    src = _docx_toc_with_bookmark(tmp_path)
    ws = tmp_path / "ws"
    extract_file(src, ws, overwrite=True)
    blocks = ContentBlocksFile.model_validate_json((ws / "content.blocks.json").read_text())
    content_md = (ws / "content.md").read_text(encoding="utf-8")
    tree = OutlineTree(
        strategy="toc",
        nodes=[
            OutlineNode(
                node_id="n1",
                title="一、 服务方案",
                level=1,
                parent_id=None,
                sort_order=0,
                anchor=Anchor(block_index=1),
                source_refs=["toc_bookmark:_TocMissing"],
            )
        ],
    )
    original_index = tree.nodes[0].anchor.block_index
    enriched = apply_toc_bookmark_anchors(tree, src, blocks, content_md=content_md)
    assert enriched.nodes[0].anchor.block_index == original_index
